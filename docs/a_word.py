"""Convierte la lista de flowables de reportlab de un informe en un .docx editable (python-docx).
Uso: python docs/a_word.py   -> genera las versiones Word de los tres informes."""
import html
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.platypus import Image, Paragraph, Table

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "docs"))

AZUL = RGBColor(0x1F, 0x4E, 0x8C)
GRIS = RGBColor(0x5A, 0x59, 0x55)
TOKEN = re.compile(r"(<br\s*/?>|</?b>|</?i>|</?sub>|</?sup>)")


def _runs(par, markup, size=None, color=None, italic=False):
    b = i = sub = sup = False
    for tok in TOKEN.split(markup):
        if not tok:
            continue
        if re.fullmatch(r"<br\s*/?>", tok):
            par.add_run().add_break()
        elif tok in ("<b>", "</b>"):
            b = tok == "<b>"
        elif tok in ("<i>", "</i>"):
            i = tok == "<i>"
        elif tok in ("<sub>", "</sub>"):
            sub = tok == "<sub>"
        elif tok in ("<sup>", "</sup>"):
            sup = tok == "<sup>"
        else:
            r = par.add_run(html.unescape(re.sub(r"<[^>]+>", "", tok)))
            r.bold, r.italic = b, (i or italic)
            r.font.subscript, r.font.superscript = sub, sup
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color


def _sombra(celda, hex_):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def convertir(story, salida):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for h, size in (("Heading 1", 13), ("Heading 2", 11), ("Title", 18)):
        doc.styles[h].font.color.rgb = AZUL
        doc.styles[h].font.size = Pt(size)

    for fl in story:
        if isinstance(fl, Paragraph):
            nombre, txt = fl.style.name, fl.text
            if nombre == "titulo":
                p = doc.add_paragraph(style="Title")
                _runs(p, txt)
            elif nombre == "h1":
                p = doc.add_paragraph(style="Heading 1")
                _runs(p, txt)
            elif nombre == "h2":
                p = doc.add_paragraph(style="Heading 2")
                _runs(p, txt)
            elif nombre == "li":
                p = doc.add_paragraph(style="List Bullet")
                _runs(p, txt)
            elif nombre == "sub":
                p = doc.add_paragraph()
                _runs(p, txt, size=9, color=GRIS)
            elif nombre == "cap":
                p = doc.add_paragraph()
                _runs(p, txt, size=8.5, color=GRIS, italic=True)
            elif nombre == "ref":
                p = doc.add_paragraph()
                _runs(p, txt, size=8.5)
            else:
                p = doc.add_paragraph()
                _runs(p, txt)
        elif isinstance(fl, Image):
            ancho = fl.drawWidth / 28.3465  # puntos -> cm
            doc.add_picture(fl.filename if isinstance(fl.filename, str) else str(fl.filename), width=Cm(min(ancho, 17.5)))
        elif isinstance(fl, Table):
            datos = fl._cellvalues
            anchos = fl._colWidths
            t = doc.add_table(rows=len(datos), cols=len(datos[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for j, col in enumerate(t.columns):
                col.width = Cm(anchos[j] / 28.3465)
            for i, fila in enumerate(datos):
                for j, c in enumerate(fila):
                    celda = t.cell(i, j)
                    celda.width = Cm(anchos[j] / 28.3465)
                    p = celda.paragraphs[0]
                    if isinstance(c, (list, tuple)):
                        txt = "<br/>".join(getattr(x, "text", str(x)) for x in c)
                    else:
                        txt = getattr(c, "text", str(c))
                    if i == 0:
                        _runs(p, txt, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
                        for r in p.runs:
                            r.bold = True
                        _sombra(celda, "1F4E8C")
                    else:
                        _runs(p, txt, size=8.5)
            doc.add_paragraph()
    doc.save(salida)
    return salida


if __name__ == "__main__":
    import act1_plan
    import act2_reporte
    import informe_final

    for mod, nombre in [(act1_plan, "Actividad1_Plan_Experimental_Grupo7.docx"),
                        (act2_reporte, "Actividad2_Reporte_Tecnico_Grupo7.docx"),
                        (informe_final, "Actividad3_Informe_Final_Grupo7.docx")]:
        story = getattr(mod, "final", None) or mod.s
        print(convertir(story, ROOT / "docs" / nombre))
